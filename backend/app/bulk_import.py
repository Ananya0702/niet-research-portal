"""Enterprise bulk import service for Schools, Departments and Faculty.

Handles CSV / XLSX parsing, row-wise validation, import modes
(insert | update | upsert), MySQL transactions and audit logging.
"""

import csv
import io
import json
import re
import secrets
import string

import openpyxl
from sqlalchemy.orm import Session

from app.models import School, Department, Faculty, SecurityAuditLog
from app.security import get_password_hash
from app.password_policy import policy_errors

ENTITIES = ("schools", "departments", "faculty")

CANONICAL_COLUMNS = {
    "schools":      ["school_name"],
    "departments":  ["school_name", "dept_code", "dept_name", "hod_faculty_id", "established_year"],
    "faculty":      ["name", "email", "phone", "branch", "school_name", "dept_code", "dept_name",
                     "employee_id", "role", "designation", "is_phd", "joining_year",
                     "google_scholar_id", "orcid_id", "scopus_author_id"],
}

SAMPLE_ROWS = {
    "schools":     [{"school_name": "School of Engineering and Technology"}],
    "departments": [{"school_name": "School of Engineering and Technology", "dept_code": "CSE",
                     "dept_name": "Computer Science and Engineering", "hod_faculty_id": "",
                     "established_year": 2004}],
    "faculty":     [{"name": "Dr. Jane Doe", "email": "jane.doe@niet.co.in", "phone": "9876543210",
                     "branch": "CSE", "school_name": "School of Engineering and Technology",
                     "dept_code": "", "dept_name": "", "employee_id": "",
                     "role": "faculty", "designation": "Assistant Professor",
                     "is_phd": "true", "joining_year": 2019, "google_scholar_id": "",
                     "orcid_id": "", "scopus_author_id": ""}],
}

ROLE_ALLOWED = {"faculty", "hod", "dean_rd"}
DEFAULT_DESIGNATION = "Assistant Professor"

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

_ALIASES = {
    "schoolname":   "school_name",
    "school":       "school_name",
    "schoolid":     "school_id",
    "deptcode":     "dept_code",
    "departmentcode": "dept_code",
    "code":         "dept_code",
    "deptname":     "dept_name",
    "department":   "dept_name",
    "departmentname": "dept_name",
    "dept":         "dept_name",
    "hodfacultyid": "hod_faculty_id",
    "hodid":        "hod_faculty_id",
    "hod":          "hod_faculty_id",
    "establishedyear": "established_year",
    "yearestablished": "established_year",
    "facultyname":  "name",
    "fullname":     "name",
    "teachername":  "name",
    "name":         "name",
    "email":        "email",
    "emailaddress": "email",
    "phone":        "phone",
    "phonenumber":  "phone",
    "phone_no":     "phone",
    "mobile":       "phone",
    "mobilenumber": "phone",
    "mobile_no":    "phone",
    "contact":      "phone",
    "contactnumber": "phone",
    "branch":       "dept_code",
    "branchcode":   "dept_code",
    "branchname":   "dept_name",
    "employeeid":   "employee_id",
    "employee_code":"employee_id",
    "empcode":      "employee_id",
    "empid":        "employee_id",
    "teacherid":    "employee_id",
    "deptid":       "dept_id",
    "role":         "role",
    "designation":  "designation",
    "design":       "designation",
    "isphd":        "is_phd",
    "phd":          "is_phd",
    "joiningyear":  "joining_year",
    "yearofjoining": "joining_year",
    "joined":       "joining_year",
    "googlescholarid": "google_scholar_id",
    "scholarid":    "google_scholar_id",
    "googlescholar": "google_scholar_id",
    "orcidid":      "orcid_id",
    "orcid":        "orcid_id",
    "scopusauthorid": "scopus_author_id",
    "scopusid":     "scopus_author_id",
    "scopus":       "scopus_author_id",
}


def _norm_key(key):
    return re.sub(r"[\s_\-\.\(\)]+", "", str(key or "").strip().lower())


def normalize_headers(raw_headers):
    """Map arbitrary file headers to canonical column names."""
    mapping = {}
    for h in raw_headers:
        n = _norm_key(h)
        if n in _ALIASES:
            mapping[h] = _ALIASES[n]
        elif n in CANONICAL_COLUMNS["schools"] or n in CANONICAL_COLUMNS["departments"] or n in CANONICAL_COLUMNS["faculty"]:
            mapping[h] = n
        else:
            mapping[h] = n
    return mapping


def normalize_rows(rows):
    """Normalize dict keys in JSON rows (wizard path) via the alias map."""
    out = []
    for row in rows or []:
        mapping = normalize_headers(list(row.keys()))
        out.append({mapping.get(k, k): v for k, v in row.items()})
    return out


def parse_rows(filename, content):
    """Parse CSV or XLSX bytes into a list of row dicts (header-normalized)."""
    name = (filename or "upload.csv").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(content)
    return _parse_csv(content)


def _clean_value(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else ""
    s = str(v).strip()
    if isinstance(v, float) and v.is_integer():
        s = str(int(v))
    return s


def _parse_csv(content):
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [r for r in reader if any(str(c).strip() for c in r)]
    if not rows:
        return []
    headers = rows[0]
    mapping = normalize_headers(headers)
    data = []
    for line in rows[1:]:
        row = {}
        for i, h in enumerate(headers):
            val = _clean_value(line[i]) if i < len(line) else ""
            row[mapping.get(h, h)] = val
        data.append(row)
    return data


def _parse_xlsx(content):
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    headers = None
    data = []
    for row in it:
        if headers is None:
            headers = [_clean_value(c) for c in row]
            mapping = normalize_headers(headers)
            continue
        if not any(c not in (None, "") for c in row):
            continue
        rec = {}
        for i, h in enumerate(headers):
            val = _clean_value(row[i]) if i < len(row) else ""
            rec[mapping.get(h, h)] = val
        data.append(rec)
    wb.close()
    return data


def _to_int(v):
    s = _clean_value(v)
    if not s:
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def _to_bool(v):
    s = _clean_value(v).lower()
    return s in ("1", "true", "yes", "y", "phd", "on")


def _gen_temp_password():
    alphabet = string.ascii_letters + string.digits + "!@#$%^&*"
    while True:
        pwd = "Niet@" + "".join(secrets.choice(alphabet) for _ in range(8))
        if not policy_errors(pwd):
            return pwd


def _audit(db, entity, mode, actor, summary, row_errors):
    detail = {
        "entity": entity,
        "mode": mode,
        "summary": summary,
        "errors": row_errors[:20],
    }
    db.add(SecurityAuditLog(
        action="bulk_import",
        method="import",
        status="success",
        detail=json.dumps(detail)[:500],
        ip_address=actor,
        device=entity[:50],
    ))
    db.commit()


# ── Schools ──────────────────────────────────────────────────────────────────
def _validate_school_row(db, row, mode, existing_names, seen):
    errors = []
    summary_bucket = None
    name = _clean_value(row.get("school_name"))
    if not name:
        errors.append("school_name is required")
        summary_bucket = "missing_fields"
        return name, errors, summary_bucket, None

    key = name.lower()
    if key in seen:
        errors.append("Duplicate school in file: " + name)
        summary_bucket = "duplicates"
        return name, errors, summary_bucket, None
    seen.add(key)

    existing = existing_names.get(key)
    if existing:
        if mode == "insert":
            errors.append("School already exists: " + name)
            summary_bucket = "duplicates"
            return name, errors, summary_bucket, None
        action = "update"
    else:
        if mode == "update":
            errors.append("School not found for update: " + name)
            summary_bucket = "missing_fields"
            return name, errors, summary_bucket, None
        action = "insert"
    return name, errors, summary_bucket, action


def _import_schools(db, rows, mode, actor, dry_run):
    existing = {s.school_name.lower(): s for s in db.query(School).all()}
    seen = set()
    valid = []
    row_errors = []
    summary = {"to_insert": 0, "to_update": 0, "duplicates": 0, "missing_fields": 0,
               "invalid_email": 0, "unknown_school": 0, "unknown_department": 0, "invalid_ids": 0}

    for i, row in enumerate(rows, start=2):
        name, errs, bucket, action = _validate_school_row(db, row, mode, existing, seen)
        if errs:
            row_errors.append({"row": i, "errors": errs, "name": name or f"Row {i}"})
            if bucket:
                summary[bucket] += 1
        else:
            valid.append({"row": i, "name": name, "action": action})
            summary["to_insert" if action == "insert" else "to_update"] += 1

    results = {"inserted": 0, "updated": 0, "skipped": 0, "failed": 0}
    if not dry_run and valid:
        for item in valid:
            try:
                with db.begin_nested():
                    if item["action"] == "insert":
                        db.add(School(school_name=item["name"]))
                        results["inserted"] += 1
                    else:
                        existing[item["name"].lower()].school_name = item["name"]
                        results["updated"] += 1
            except Exception:
                db.rollback()
                results["failed"] += 1
        try:
            db.commit()
            try:
                _audit(db, "schools", mode, actor, summary, row_errors)
            except Exception:
                db.rollback()
        except Exception:
            db.rollback()
            results["failed"] = results["inserted"] + results["updated"]
            results["inserted"] = results["updated"] = 0

    return {"entity": "schools", "mode": mode, "dry_run": dry_run,
            "total_rows": len(rows), "valid_rows": len(valid), "invalid_rows": len(row_errors),
            "summary": summary, "row_errors": row_errors[:200], "results": results}


# ── Departments ──────────────────────────────────────────────────────────────
def _resolve_school(db, row):
    schools = {s.school_name.lower(): s for s in db.query(School).all()}
    sid = _to_int(row.get("school_id"))
    if sid:
        s = db.query(School).filter(School.school_id == sid).first()
        if s:
            return s, None
        return None, "Unknown school_id"
    sname = _clean_value(row.get("school_name"))
    if sname:
        s = schools.get(sname.lower())
        if s:
            return s, None
        return None, "Unknown school: " + sname
    return None, "school_name is required"


def _validate_dept_row(db, row, mode, existing_codes, seen):
    errors = []
    bucket = None
    action = None

    code = _clean_value(row.get("dept_code"))
    name = _clean_value(row.get("dept_name"))
    if not code:
        errors.append("dept_code is required")
        bucket = "missing_fields"
    if not name:
        errors.append("dept_name is required")
        bucket = bucket or "missing_fields"
    if not code:
        return errors, bucket, action

    key = code.lower()
    if key in seen:
        errors.append("Duplicate dept_code in file: " + code)
        bucket = "duplicates"
        return errors, bucket, action
    seen.add(key)

    school, err = _resolve_school(db, row)
    if err:
        errors.append(err)
        bucket = bucket or "unknown_school"
    if errors:
        return errors, bucket, action

    hod = _to_int(row.get("hod_faculty_id"))
    year = _to_int(row.get("established_year"))

    existing = existing_codes.get(key)
    if existing:
        if mode == "insert":
            errors.append("Department already exists: " + code)
            bucket = "duplicates"
            return errors, bucket, action
        action = "update"
    else:
        if mode == "update":
            errors.append("Department not found for update: " + code)
            bucket = bucket or "missing_fields"
            return errors, bucket, action
        action = "insert"
    return errors, bucket, action, {"code": code, "name": name, "school": school,
                                    "hod": hod, "year": year, "existing": existing if action == "update" else None}


def _import_departments(db, rows, mode, actor, dry_run):
    existing = {d.dept_code.lower(): d for d in db.query(Department).all()}
    seen = set()
    valid = []
    row_errors = []
    summary = {"to_insert": 0, "to_update": 0, "duplicates": 0, "missing_fields": 0,
               "invalid_email": 0, "unknown_school": 0, "unknown_department": 0, "invalid_ids": 0}

    for i, row in enumerate(rows, start=2):
        result = _validate_dept_row(db, row, mode, existing, seen)
        if len(result) == 3:
            errors, bucket, action = result
            code = _clean_value(row.get("dept_code")) or f"Row {i}"
            row_errors.append({"row": i, "errors": errors, "name": code})
            if bucket:
                summary[bucket] += 1
            continue
        errors, bucket, action, data = result
        if errors:
            row_errors.append({"row": i, "errors": errors, "name": data["code"]})
            if bucket:
                summary[bucket] += 1
        else:
            valid.append({"row": i, "action": action, "data": data})
            summary["to_insert" if action == "insert" else "to_update"] += 1

    results = {"inserted": 0, "updated": 0, "skipped": 0, "failed": 0}
    if not dry_run and valid:
        for item in valid:
            d = item["data"]
            try:
                with db.begin_nested():
                    if item["action"] == "insert":
                        db.add(Department(school_id=d["school"].school_id, dept_code=d["code"],
                                          dept_name=d["name"], hod_faculty_id=d["hod"],
                                          established_year=d["year"]))
                        results["inserted"] += 1
                    else:
                        dept = d["existing"]
                        dept.school_id = d["school"].school_id
                        dept.dept_name = d["name"]
                        dept.hod_faculty_id = d["hod"]
                        dept.established_year = d["year"]
                        results["updated"] += 1
            except Exception:
                db.rollback()
                results["failed"] += 1
        try:
            db.commit()
            try:
                _audit(db, "departments", mode, actor, summary, row_errors)
            except Exception:
                db.rollback()
        except Exception:
            db.rollback()
            results["failed"] = results["inserted"] + results["updated"]
            results["inserted"] = results["updated"] = 0

    return {"entity": "departments", "mode": mode, "dry_run": dry_run,
            "total_rows": len(rows), "valid_rows": len(valid), "invalid_rows": len(row_errors),
            "summary": summary, "row_errors": row_errors[:200], "results": results}


# ── Faculty ──────────────────────────────────────────────────────────────────
def _resolve_department(db, row):
    did = _to_int(row.get("dept_id"))
    if did:
        d = db.query(Department).filter(Department.dept_id == did).first()
        if d:
            return d, None
        return None, "Unknown dept_id: " + str(did)
    depts = {d.dept_code.lower(): d for d in db.query(Department).all()}
    by_name = {d.dept_name.lower(): d for d in db.query(Department).all()}
    code = _clean_value(row.get("dept_code"))
    if code:
        d = depts.get(code.lower())
        if d:
            return d, None
        return None, "Unknown department: " + code
    dname = _clean_value(row.get("dept_name"))
    if dname:
        d = by_name.get(dname.lower())
        if d:
            return d, None
        return None, "Unknown department: " + dname
    return None, "dept_id, dept_code or dept_name is required"


def _resolve_faculty_school(db, row):
    sid = _to_int(row.get("school_id"))
    if sid:
        s = db.query(School).filter(School.school_id == sid).first()
        if s:
            return s, None
        return None, "Unknown school_id: " + str(sid)
    sname = _clean_value(row.get("school_name"))
    if not sname:
        return None, None
    schools = {s.school_name.lower(): s for s in db.query(School).all()}
    s = schools.get(sname.lower())
    if s:
        return s, None
    return None, "Unknown school: " + sname


def _gen_employee_id(db):
    """Auto-generate a unique employee / teacher id like NIET-FAC-0001."""
    prefix = "NIET-FAC-"
    max_seq = 0
    for f in db.query(Faculty).all():
        eid = (f.employee_id or "").strip()
        if eid.startswith(prefix):
            try:
                max_seq = max(max_seq, int(eid[len(prefix):]))
            except ValueError:
                pass
    return prefix + str(max_seq + 1).zfill(4)


def _validate_faculty_row(db, row, mode, existing_emails, seen):
    errors = []
    bucket = None
    action = None

    name = _clean_value(row.get("name"))
    email = _clean_value(row.get("email")).lower()
    phone = _clean_value(row.get("phone"))
    if not name:
        errors.append("name is required")
        bucket = "missing_fields"
    if not email:
        errors.append("email is required")
        bucket = bucket or "missing_fields"
    elif not EMAIL_RE.match(email):
        errors.append("Invalid email format: " + email)
        bucket = "invalid_email"
    if errors:
        return errors, bucket, action

    if email in seen:
        errors.append("Duplicate email in file: " + email)
        bucket = "duplicates"
        return errors, bucket, action
    seen.add(email)

    dept, err = _resolve_department(db, row)
    if err:
        errors.append(err)
        bucket = bucket or "unknown_department"

    school, err2 = _resolve_faculty_school(db, row)
    if err2:
        errors.append(err2)
        bucket = bucket or "unknown_school"

    if errors:
        return errors, bucket, action

    if dept.school_id is not None and school is not None and dept.school_id != school.school_id:
        errors.append(f"School '{school.school_name}' does not match department '{dept.dept_code}'")
        bucket = bucket or "unknown_school"
        return errors, bucket, action

    role = _clean_value(row.get("role")) or "faculty"
    role = role.lower().replace(" ", "_")
    if role == "dean":
        role = "dean_rd"
    if role not in ROLE_ALLOWED:
        errors.append("Invalid role: " + role)
        bucket = "invalid_ids"
        return errors, bucket, action

    designation = _clean_value(row.get("designation")) or DEFAULT_DESIGNATION
    is_phd = _to_bool(row.get("is_phd"))
    joining_year = _to_int(row.get("joining_year"))
    gscholar = _clean_value(row.get("google_scholar_id"))
    orcid = _clean_value(row.get("orcid_id"))
    scopus = _clean_value(row.get("scopus_author_id"))
    employee_id = _clean_value(row.get("employee_id"))

    existing = existing_emails.get(email)
    if existing:
        if mode == "insert":
            errors.append("Email already registered: " + email)
            bucket = "duplicates"
            return errors, bucket, action
        action = "update"
    else:
        if mode == "update":
            errors.append("Faculty not found for update: " + email)
            bucket = bucket or "missing_fields"
            return errors, bucket, action
        action = "insert"
    return errors, bucket, action, {"name": name, "email": email, "dept": dept, "school": school,
                                    "role": role, "phone": phone,
                                    "designation": designation, "is_phd": is_phd,
                                    "joining_year": joining_year, "gscholar": gscholar,
                                    "orcid": orcid, "scopus": scopus,
                                    "employee_id": employee_id,
                                    "existing": existing if action == "update" else None,
                                    "temp_password": None}


def _import_faculty(db, rows, mode, actor, dry_run):
    existing = {f.email.lower(): f for f in db.query(Faculty).all()}
    seen = set()
    valid = []
    row_errors = []
    summary = {"to_insert": 0, "to_update": 0, "duplicates": 0, "missing_fields": 0,
               "invalid_email": 0, "unknown_school": 0, "unknown_department": 0, "invalid_ids": 0}

    for i, row in enumerate(rows, start=2):
        result = _validate_faculty_row(db, row, mode, existing, seen)
        if len(result) == 3:
            errors, bucket, action = result
            email = _clean_value(row.get("email")) or f"Row {i}"
            row_errors.append({"row": i, "errors": errors, "name": email})
            if bucket:
                summary[bucket] += 1
            continue
        errors, bucket, action, data = result
        if errors:
            row_errors.append({"row": i, "errors": errors, "name": data.get("email") or f"Row {i}"})
            if bucket:
                summary[bucket] += 1
        else:
            if action == "insert":
                data["temp_password"] = _gen_temp_password()
            valid.append({"row": i, "action": action, "data": data})
            summary["to_insert" if action == "insert" else "to_update"] += 1

    results = {"inserted": 0, "updated": 0, "skipped": 0, "failed": 0}
    temp_passwords = {}
    generated_ids = {}
    if not dry_run and valid:
        for item in valid:
            d = item["data"]
            try:
                with db.begin_nested():
                    if item["action"] == "insert":
                        eid = d["employee_id"] or _gen_employee_id(db)
                        generated_ids[item["row"]] = {"email": d["email"], "employee_id": eid}
                        fac = Faculty(
                            name=d["name"], email=d["email"],
                            password_hash=get_password_hash(d["temp_password"]),
                            role=d["role"], designation=d["designation"],
                            dept_id=d["dept"].dept_id, school_id=d["school"].school_id if d["school"] else d["dept"].school_id,
                            is_phd=d["is_phd"], joining_year=d["joining_year"],
                            google_scholar_id=d["gscholar"] or None,
                            orcid_id=d["orcid"] or None, scopus_author_id=d["scopus"] or None,
                            phone=d["phone"] or None, employee_id=eid,
                            is_active=True, force_password_reset=True,
                            h_index=0, i10_index=0, total_publications=0, total_citations=0,
                        )
                        db.add(fac)
                        db.flush()
                        from app.patent_links import link_saved_patents_to_new_faculty
                        link_saved_patents_to_new_faculty(db, fac)
                        temp_passwords[item["row"]] = {"email": d["email"], "password": d["temp_password"]}
                        results["inserted"] += 1
                    else:
                        f = d["existing"]
                        f.name = d["name"]
                        f.role = d["role"]
                        f.designation = d["designation"]
                        f.dept_id = d["dept"].dept_id
                        f.school_id = d["school"].school_id if d["school"] else d["dept"].school_id
                        f.is_phd = d["is_phd"]
                        f.joining_year = d["joining_year"]
                        f.google_scholar_id = d["gscholar"] or f.google_scholar_id
                        f.orcid_id = d["orcid"] or f.orcid_id
                        f.scopus_author_id = d["scopus"] or f.scopus_author_id
                        if d["phone"]:
                            f.phone = d["phone"]
                        if d["employee_id"]:
                            f.employee_id = d["employee_id"]
                        results["updated"] += 1
            except Exception:
                db.rollback()
                results["failed"] += 1
        try:
            db.commit()
            try:
                _audit(db, "faculty", mode, actor, summary, row_errors)
            except Exception:
                db.rollback()
        except Exception:
            db.rollback()
            results["failed"] = results["inserted"] + results["updated"]
            results["inserted"] = results["updated"] = 0

    return {"entity": "faculty", "mode": mode, "dry_run": dry_run,
            "total_rows": len(rows), "valid_rows": len(valid), "invalid_rows": len(row_errors),
            "summary": summary, "row_errors": row_errors[:200], "results": results,
            "temp_passwords": temp_passwords, "generated_ids": generated_ids}


# ── Public entry ─────────────────────────────────────────────────────────────
def run_import(db, entity, rows, mode="upsert", dry_run=False, actor="admin"):
    if entity not in ENTITIES:
        raise ValueError("Unsupported import entity: " + entity)
    if mode not in ("insert", "update", "upsert"):
        raise ValueError("mode must be insert, update or upsert")
    rows = normalize_rows(rows or [])
    if entity == "schools":
        return _import_schools(db, rows, mode, actor, dry_run)
    if entity == "departments":
        return _import_departments(db, rows, mode, actor, dry_run)
    return _import_faculty(db, rows, mode, actor, dry_run)


def generate_template_csv(entity, with_sample=False):
    if entity not in ENTITIES:
        raise ValueError("Unsupported entity")
    headers = CANONICAL_COLUMNS[entity]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    if with_sample:
        for row in SAMPLE_ROWS[entity]:
            w.writerow([row.get(h, "") for h in headers])
    return buf.getvalue()


def generate_template_xlsx(entity, with_sample=False):
    headers = CANONICAL_COLUMNS[entity]
    rows = [headers]
    if with_sample:
        for row in SAMPLE_ROWS[entity]:
            rows.append([row.get(h, "") for h in headers])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = entity
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
